import sys
sys.path.append('../')
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from lib.SusiiProductLoader import SusiiProductLoader

class HomologarProducts:
    def exportarProductos(self, businessId, sede):
        print(f"\nDescargando productos de {sede}...")
        loader = SusiiProductLoader(businessId)
        productos = loader.downloadProducts(includeDisable=True)
        datos = []
        for producto in productos.values():
            datos.append({
                "CODIGO": producto.getCode(),
                "NOMBRE": producto.getName(),
                "PRECIO": producto.getPrice(),
                "DISABLE": producto.isDisable(),
                "GENERICO": producto.getGenerico(),
                "SEDE": sede
            })
        df = pd.DataFrame(datos)
        nombreArchivo = f"Productos_{sede}.xlsx"
        df.to_excel(nombreArchivo, index=False)
        print(f"{nombreArchivo} generado, con ({len(df)} productos).")
        return df
    
    def compararProductos(self, dfRetamas, dfCobian):
        print("\nComparando productos...")
        retamas = dfRetamas.set_index("CODIGO")
        cobian = dfCobian.set_index("CODIGO")
        wb = Workbook()
        wsCoin = wb.active
        wsCoin.title = "Diferencias Por Codigo"
        wsCoin.append([
            "CODIGO",
            "NOMBRE RETAMAS",
            "PRECIO RETAMAS",
            "NOMBRE COBIAN",
            "PRECIO COBIAN",
            "ESTADO"
        ])
        wsSoloRet = wb.create_sheet("Solo Retamas")
        wsSoloRet.append([
            "CODIGO",
            "NOMBRE",
            "PRECIO"
        ])
        wsSoloCob = wb.create_sheet("Solo Cobian")
        wsSoloCob.append([
            "CODIGO",
            "NOMBRE",
            "PRECIO"
        ])
        wsDisable = wb.create_sheet("Productos Deshabilitados")
        wsDisable.append([
            "CODIGO",
            "NOMBRE RETAMAS",
            "ESTADO RETAMAS",
            "NOMBRE COBIAN",
            "ESTADO COBIAN"
        ])
        wsGold = wb.create_sheet("GOLD Desincronizados")
        wsGold.append([
            "CODIGO",
            "NOMBRE RETAMAS",
            "GENERICO RETAMAS",
            "NOMBRE COBIAN",
            "GENERICO COBIAN",
            "ESTADO"
        ])
        coincidencias = 0
        disableCount = 0
        for codigo in retamas.index:
            if codigo in cobian.index:
                r = retamas.loc[codigo]
                c = cobian.loc[codigo]
                # VALIDAR GOLD
                genericoRetamas = r["GENERICO"]
                genericoCobian = c["GENERICO"]
                # Solo interesa cuando exactamente una sede tiene GENERICO = 2
                if (genericoRetamas == 2) != (genericoCobian == 2):
                    if genericoRetamas == 2:
                        estado = "Cobian no es gold"
                    else:
                        estado = "Retamas no es gold"
                    wsGold.append([
                        codigo,
                        r["NOMBRE"],
                        genericoRetamas,
                        c["NOMBRE"],
                        genericoCobian,
                        estado
                    ])
                if r["DISABLE"] and c["DISABLE"]:
                    continue
                if r["DISABLE"] != c["DISABLE"]:
                    wsDisable.append([
                        codigo,
                        r["NOMBRE"], "DISABLE" if r["DISABLE"] else "ACTIVO",
                        c["NOMBRE"], "DISABLE" if c["DISABLE"] else "ACTIVO"
                    ])
                    disableCount +=1
                    continue
                nombreIgual = (r["NOMBRE"] == c["NOMBRE"])
                precioIgual = (r["PRECIO"] == c["PRECIO"])
                if nombreIgual and precioIgual:
                    continue
                if not nombreIgual and not precioIgual:
                    estado = "Nombre y precio diferentes"
                elif not nombreIgual:
                    estado = "Nombre diferente"
                else:
                    estado = "Precio diferente"
                wsCoin.append([
                    codigo,
                    r["NOMBRE"],
                    r["PRECIO"],
                    c["NOMBRE"],
                    c["PRECIO"],
                    estado
                ])
                coincidencias += 1
            else:
                r = retamas.loc[codigo]
                if not r["DISABLE"]:
                    wsSoloRet.append([
                        codigo,
                        r["NOMBRE"],
                        r["PRECIO"]
                    ])
        for codigo in cobian.index:
            if codigo not in retamas.index:
                c = cobian.loc[codigo]
                if not c["DISABLE"]:
                    wsSoloCob.append([
                        codigo,
                        c["NOMBRE"],
                        c["PRECIO"]
                    ])
        fecha = datetime.now().strftime("%Y-%m-%d")
        nombreArchivo = f"out/Homologacion_Productos_{fecha}.xlsx"
        wb.save(nombreArchivo)
        print(f"Diferencias encontradas: {coincidencias}")
        print(f"Productos deshabilitados: {disableCount}")
        print(f"Solo Retamas: {wsSoloRet.max_row - 1}")
        print(f"Solo Cobian : {wsSoloCob.max_row - 1}")
        print(f"Archivo generado: {nombreArchivo}")

    def run(self):
        print("========================================")
        print("  HOMOLOGADOR DE PRODUCTOS QANTUFARMA   ")
        print("========================================")
        dfRetamas = self.exportarProductos(8132, sede="Retamas")
        dfCobian = self.exportarProductos(5053, sede="Cobian")
        self.compararProductos(dfRetamas, dfCobian)
        print("\nDescarga finalizada correctamente.")
        
if __name__ == "__main__":
    HomologarProducts().run()